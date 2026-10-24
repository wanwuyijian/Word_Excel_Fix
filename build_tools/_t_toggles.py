"""逐项验证 8 个交互开关"确实有效"（不是只连了线）。

做法：每个开关单独打开（其余全关），对比输出与"全关闭"基线，
找出该开关应当产生差异的观测点；有差异 = 生效。
"""
from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from docx import Document  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Font as XLFont, PatternFill  # noqa: E402

from wxfixer.core import optimize_file  # noqa: E402
from wxfixer.core.runner import OptimizeOptions  # noqa: E402
from wxfixer.templates import get_template  # noqa: E402
from wxfixer.templates.schema import TOGGLE_META, ToggleOptions  # noqa: E402

PASS: list[str] = []
FAIL: list[str] = []


def check(cond: bool, msg: str) -> None:
    (PASS if cond else FAIL).append(msg)
    print(("  OK   " if cond else "  FAIL ") + msg)


def png_bytes(w: int, h: int, rgb=(60, 110, 190)) -> bytes:
    import struct
    import zlib

    raw = b""
    row = bytes(rgb) * w
    for _ in range(h):
        raw += b"\x00" + row

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw, 6)) + chunk(b"IEND", b""))


# ---------------------------------------------------------------------------
def build_word(path: str, png: str) -> None:
    doc = Document()
    doc.add_paragraph("设备采购技术方案")
    doc.add_paragraph("本方案用于指导本次设备采购与安装调试工作。")
    doc.add_paragraph("一、项目背景")
    p = doc.add_paragraph("本项目建设内容包括机房改造、网络设备采购与系统集成调试。")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.name = "Times New Roman"
    doc.add_paragraph("1.1 建设目标")
    doc.add_paragraph("  众所周知， 该 设备 性能 很好。  ")
    doc.add_paragraph().add_run().add_picture(png, width=Cm(22))
    doc.add_paragraph("图 1  现场示意图")
    doc.add_paragraph("二、建设内容")
    doc.add_paragraph("包含机房装修、供配电改造、机柜与网络设备安装。")
    t = doc.add_table(rows=3, cols=3)
    for i, row in enumerate((("序号", "名称", "数量"), ("1", "服务器", "2"),
                             ("2", "交换机", "4"))):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    doc.save(path)


def build_excel(path: str, png: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "清单"
    ws.append(["序号", "名称", "备注"])
    ws.append([1, "服务器", "  众所周知， 该 设备 性能 很好。  "])
    ws.append([2, "交换机", "基本上,运行稳定。"])
    for row in ws.iter_rows(min_row=1, max_row=3, max_col=3):
        for c in row:
            c.font = XLFont(name="Times New Roman", size=9, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    img = XLImage(png)
    img.width, img.height = 1600, 1000
    ws.add_image(img, "E3")
    wb.save(path)


# ---------------------------------------------------------------------------
def word_probe(path: str) -> dict:
    doc = Document(path)
    first_run = None
    for p in doc.paragraphs:
        if p.runs and p.text.strip():
            first_run = p.runs[0]
            break
    sizes = set()
    colors = set()
    names = set()
    shades = set()
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size:
                sizes.add(round(r.font.size.pt, 1))
            if r.font.name:
                names.add(r.font.name)
            try:
                if r.font.color and r.font.color.rgb:
                    colors.add(str(r.font.color.rgb))
            except Exception:
                pass
        shd = p._p.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
        if shd is not None:
            s = shd.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
            if s is not None:
                shades.add(s.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill"))
    cell_shades = set()
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                tcpr = c._tc.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr")
                if tcpr is not None:
                    s = tcpr.find(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
                    if s is not None:
                        cell_shades.add(s.get(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill"))
    breaks = sum(1 for p in doc.paragraphs if p.paragraph_format.page_break_before)
    widths = [round(s.width.cm, 2) for s in doc.inline_shapes]
    texts = [p.text for p in doc.paragraphs]
    return {
        "sizes": sizes, "colors": colors, "names": names,
        "para_shades": shades, "cell_shades": cell_shades,
        "breaks": breaks, "widths": widths, "texts": texts,
        "sample_run": (first_run.font.name if first_run else None,
                       str(first_run.font.size) if first_run else None),
    }


def excel_probe(path: str) -> dict:
    wb = load_workbook(path)
    ws = wb["清单"]
    names, sizes, colors, fills = set(), set(), set(), set()
    for row in ws.iter_rows(min_row=1, max_row=3, max_col=3):
        for c in row:
            if c.font:
                if c.font.name:
                    names.add(c.font.name)
                if c.font.sz:
                    sizes.add(round(float(c.font.sz), 1))
                try:
                    if c.font.color and c.font.color.rgb:
                        colors.add(str(c.font.color.rgb))
                except Exception:
                    pass
            try:
                if c.fill and c.fill.fgColor and c.fill.fgColor.rgb:
                    fills.add(str(c.fill.fgColor.rgb))
            except Exception:
                pass
    imgs = list(getattr(ws, "_images", []))
    cx = 0
    if imgs:
        ext = getattr(imgs[0].anchor, "ext", None)
        cx = int(getattr(ext, "cx", 0) or 0)
    colw = {k: round(v.width or 0, 1) for k, v in list(ws.column_dimensions.items())[:4]}
    c2 = ws["C2"].value
    return {"names": names, "sizes": sizes, "colors": colors, "fills": fills,
            "img_cx": cx, "colw": colw, "c2": c2,
            "header_freeze": ws.freeze_panes,
            "autofilter": str(ws.auto_filter.ref)}


# ---------------------------------------------------------------------------
BASE_WORD = dict(font_family=False, font_size=False, font_color=False,
                 background=False, page_break=False, layout=False,
                 content=False, image_size=False)


def toggles_with(**kw) -> ToggleOptions:
    base = dict(BASE_WORD)
    base.update(kw)
    return ToggleOptions(**base)


def run_word(src: str, out: str) -> dict:
    print("\n[Word 开关逐项验证]")
    base_res = optimize_file(src, get_template("auto"),
                             OptimizeOptions(toggles=toggles_with()), out)
    base = word_probe(base_res.target)
    print(f"  基线(全关): 字号{base['sizes']} 字体{base['names']} "
          f"图片{base['widths']} 分页{base['breaks']}")

    # 1 修改字体
    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_family=True)), out)
    pr = word_probe(r.target)
    check(pr["sample_run"][0] != base["sample_run"][0] or pr["names"] != base["names"],
          f"字体开关：正文字体由 {base['sample_run'][0]} 变为 {pr['sample_run'][0]}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_size=True)), out)
    pr = word_probe(r.target)
    check(pr["sample_run"][1] != base["sample_run"][1],
          f"字号开关：正文字号由 {base['sample_run'][1]} 变为 {pr['sample_run'][1]}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_color=True)), out)
    pr = word_probe(r.target)
    check(pr["colors"] != base["colors"],
          f"字体颜色开关：颜色集 {base['colors']} → {pr['colors']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, background=True)), out)
    pr = word_probe(r.target)
    check(pr["cell_shades"] != base["cell_shades"],
          f"背景色开关：表头底色 {base['cell_shades']} → {pr['cell_shades']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, page_break=True)), out)
    pr = word_probe(r.target)
    check(pr["breaks"] >= base["breaks"] and pr["breaks"] > 0,
          f"分页开关：段前分页 {base['breaks']} → {pr['breaks']}")

    # 版式总开关（不带任何子项）
    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True)), out)
    pr = word_probe(r.target)
    check(pr["texts"] == base["texts"], "版式总开关：文字未变")

    # 内容
    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(content=True)), out)
    pr = word_probe(r.target)
    joined = "".join(pr["texts"])
    check("众所周知" not in joined, "内容开关：赘语已删")
    check("性能" in joined and "机房改造" in joined, "内容开关：要点保留")
    check(pr["widths"] == base["widths"], "内容开关：图片尺寸未受影响")

    # 图片
    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(image_size=True)), out)
    pr = word_probe(r.target)
    check(pr["widths"][0] < base["widths"][0],
          f"图片开关：超宽图 {base['widths'][0]}cm → {pr['widths'][0]}cm")

    # 全部关闭时输出应与源文件等价（文字与图片尺寸）
    check(base["texts"] == [p.text for p in Document(src).paragraphs],
          "全关闭：文字与源文件一致")
    check(base["widths"] == [round(s.width.cm, 2) for s in Document(src).inline_shapes],
          "全关闭：图片尺寸与源文件一致")
    return base


def _excel_fit(path: str) -> tuple:
    wb = load_workbook(path)
    ws = wb["清单"]
    try:
        fit = ws.sheet_properties.pageSetUpPr.fitToPage
    except Exception:
        fit = None
    return fit, ws.print_title_rows


def run_excel(src: str, out: str) -> None:
    print("\n[Excel 开关逐项验证]")
    base_res = optimize_file(src, get_template("auto"),
                             OptimizeOptions(toggles=toggles_with()), out)
    base = excel_probe(base_res.target)
    print(f"  基线(全关): 字号{base['sizes']} 字体{base['names']} "
          f"图片cx{base['img_cx']} 列宽{base['colw']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_family=True)), out)
    pr = excel_probe(r.target)
    check(pr["names"] != base["names"],
          f"Excel 字体开关：{base['names']} → {pr['names']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_size=True)), out)
    pr = excel_probe(r.target)
    check(pr["sizes"] != base["sizes"],
          f"Excel 字号开关：{base['sizes']} → {pr['sizes']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, font_color=True)), out)
    pr = excel_probe(r.target)
    check(pr["colors"] != base["colors"],
          f"Excel 颜色开关：{base['colors']} → {pr['colors']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True, background=True)), out)
    pr = excel_probe(r.target)
    check(pr["fills"] != base["fills"],
          f"Excel 背景开关：{base['fills']} → {pr['fills']}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(layout=True)), out)
    pr = excel_probe(r.target)
    check(pr["colw"] != base["colw"] or pr["header_freeze"] != base["header_freeze"],
          f"Excel 排版总开关：列宽/冻结 {base['colw']} → {pr['colw']}")

    r_pb = optimize_file(src, get_template("auto"),
                         OptimizeOptions(toggles=toggles_with(layout=True,
                                                              page_break=True)), out)
    r_nopb = optimize_file(src, get_template("auto"),
                           OptimizeOptions(toggles=toggles_with(layout=True,
                                                                page_break=False)), out)
    fit_pb = _excel_fit(r_pb.target)
    fit_nopb = _excel_fit(r_nopb.target)
    check(fit_pb[0] is True and fit_pb[1] is not None,
          f"Excel 分页开关开启：缩放打印={fit_pb[0]} 重复表头={fit_pb[1]}")
    check(fit_nopb[0] is not True and fit_nopb[1] is None,
          f"Excel 分页开关关闭：缩放打印={fit_nopb[0]} 重复表头={fit_nopb[1]}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(content=True)), out)
    pr = excel_probe(r.target)
    c2 = str(pr["c2"] or "")
    check("众所周知" not in c2 and "性能" in c2,
          f"Excel 内容开关：{base['c2']!r} → {c2!r}")

    r = optimize_file(src, get_template("auto"),
                      OptimizeOptions(toggles=toggles_with(image_size=True)), out)
    pr = excel_probe(r.target)
    check(0 < pr["img_cx"] < base["img_cx"],
          f"Excel 图片开关：cx {base['img_cx']} → {pr['img_cx']}")


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="wxfixer_tg_")
    work = os.path.join(tmp, "src")
    out = os.path.join(tmp, "out")
    os.makedirs(work, exist_ok=True)
    os.makedirs(out, exist_ok=True)
    png = os.path.join(work, "p.png")
    with open(png, "wb") as fh:
        fh.write(png_bytes(1400, 900))
    wsrc = os.path.join(work, "开关验证.docx")
    xsrc = os.path.join(work, "开关验证.xlsx")
    build_word(wsrc, png)
    build_excel(xsrc, png)

    run_word(wsrc, out)
    run_excel(xsrc, out)

    print("\n[元数据一致性]")
    keys = {m[0] for m in TOGGLE_META}
    check(len(keys) == 8, f"TOGGLE_META 共 {len(keys)} 项")
    check(keys == set(ToggleOptions().__dict__.keys()),
          f"TOGGLE_META 与 ToggleOptions 字段一致: {sorted(keys)}")

    print(f"\n合计：通过 {len(PASS)} / 失败 {len(FAIL)}")
    for f in FAIL:
        print("  × " + f)
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
