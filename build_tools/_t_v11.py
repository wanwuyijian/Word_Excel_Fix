"""v1.1 新开关功能验证：文字内容优化 / 图片尺寸 / 排版总开关。

用法：PYTHONPATH=src python build_tools/_t_v11.py
"""
import io
import os
import shutil
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook

from wxfixer.core import runner
from wxfixer.core.runner import OptimizeOptions
from wxfixer.templates import get_template
from wxfixer.templates.schema import ToggleOptions

OUT = os.path.join(os.path.dirname(__file__), "..", "tests", "out_v11")
os.makedirs(OUT, exist_ok=True)

# 生成一张真实 PNG（纯色块），用于图片尺寸测试
def make_png(path, w=1400, h=900, color=(60, 110, 190)):
    from PIL import Image

    Image.new("RGB", (w, h), color).save(path)


def build_word_sample(path, png):
    doc = Document()
    doc.add_heading("关于 项目 实施 的 通知", 0)
    doc.add_paragraph("众所周知，本项目 的 实施 需要 各部门 配合。")
    doc.add_paragraph("基本上,该方案可行。实际上，它能满足需求。")
    doc.add_paragraph("我们需要，尽快完成这项工作，，并 提交 报告!!!")
    doc.add_paragraph("值得一提的是，我们 的 产品 支持 多平台。")
    doc.add_paragraph("此项工作非常重要。")
    p = doc.add_paragraph()
    p.add_run().add_picture(png, width=Cm(22))     # 超出版心
    doc.add_paragraph("图 1  示意图")
    p2 = doc.add_paragraph()
    p2.add_run().add_picture(png, width=Cm(1.2))   # 图标大小
    doc.save(path)


def build_excel_sample(path, png):
    wb = Workbook()
    ws = wb.active
    ws.title = "清单"
    ws["A1"] = "设备 清单 明细表"
    ws["A2"], ws["B2"], ws["C2"] = "序号", "名称", "备注"
    ws["A3"], ws["B3"], ws["C3"] = 1, "服务器", "  众所周知，该 设备 性能 很好。  "
    ws["A4"], ws["B4"], ws["C4"] = 2, "交换机", "基本上,运行稳定。"
    from openpyxl.drawing.image import Image as XLImage

    img = XLImage(png)
    img.width, img.height = 1600, 1000
    ws.add_image(img, "E3")
    wb.save(path)


def fp(path, kind):
    if kind == "word":
        from wxfixer.core.word_optimizer import extract_text_fingerprint

        return extract_text_fingerprint(path)
    from wxfixer.core.excel_optimizer import extract_text_fingerprint as f

    return f(path)


def word_text(path):
    from wxfixer.core.word_optimizer import dump_text

    return dump_text(path)


def main():
    png = os.path.join(OUT, "_block.png")
    make_png(png)
    wsrc = os.path.join(OUT, "样例_内容优化.docx")
    xsrc = os.path.join(OUT, "样例_内容优化.xlsx")
    build_word_sample(wsrc, png)
    build_excel_sample(xsrc, png)

    tpl = get_template("auto")
    ok = 0
    bad = 0

    def check(cond, msg):
        nonlocal ok, bad
        if cond:
            ok += 1
            print("  通过", msg)
        else:
            bad += 1
            print("  失败", msg)

    # ---- 场景 1：全部关闭 ----
    print("\n[1] 全部开关关闭：应完全不动")
    r = runner.optimize_file(wsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(font_family=False, font_size=False, font_color=False,
                              background=False, page_break=False, layout=False,
                              content=False, image_size=False)), OUT)
    check(fp(r.target, "word") == fp(wsrc, "word"), "Word 文字零改动")
    check("众所周知" in word_text(r.target), "赘语未被删除（内容开关关闭）")

    # ---- 场景 2：只开内容优化 ----
    print("\n[2] 仅开启「优化文字内容」")
    r = runner.optimize_file(wsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(font_family=False, font_size=False, font_color=False,
                              background=False, page_break=False, layout=False,
                              content=True, image_size=False)), OUT)
    txt = word_text(r.target)
    check("众所周知" not in txt, "句首赘语已被删除")
    check("基本上" not in txt, "句中赘语已被删除")
    check("各部门" in txt and "多平台" in txt and "报告" in txt, "关键信息保留")
    check("!!!" not in txt, "重复标点已规整")
    check("本项目的实施需要各部门配合。" in txt, "空格已清理且语句通顺")
    check(len(txt) < len(word_text(wsrc)), "整体变短")
    check(r.stats.get("文字优化", 0) >= 3, f"统计到改动段数 {r.stats.get('文字优化')}")

    # ---- 场景 3：只开图片尺寸 ----
    print("\n[3] 仅开启「调整图片尺寸」")
    r = runner.optimize_file(wsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(font_family=False, font_size=False, font_color=False,
                              background=False, page_break=False, layout=False,
                              content=False, image_size=True)), OUT)
    d = Document(r.target)
    widths = [round(s.width.cm, 2) for s in d.inline_shapes]
    print("   图片宽度(cm)：", widths)
    check(all(w <= 16.5 for w in widths), "超宽图已缩到版心内")
    check(any(w < 3 for w in widths), "小图标保持原尺寸（未放大）")
    check(fp(r.target, "word") == fp(wsrc, "word"), "图片缩放不影响文字")

    # ---- 场景 4：内容 + 图片 + 排版 全开 ----
    print("\n[4] 三个新开关 + 原版式开关全开")
    r = runner.optimize_file(wsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(content=True)), OUT)
    check(os.path.exists(r.target), "输出文件存在")
    from wxfixer.core import ooxml_check

    issues = ooxml_check.check_document(r.target)
    check(not issues, f"OOXML 结构自检通过（问题 {len(issues)} 项）")
    if issues:
        print("   ", issues[:3])
    txt4 = word_text(r.target)
    check("众所周知" not in txt4, "内容已精简")
    d4 = Document(r.target)
    check(all(round(s.width.cm, 2) <= 16.5 for s in d4.inline_shapes), "图片已适配")

    # ---- 场景 5：Excel 内容优化 + 图片 ----
    print("\n[5] Excel：内容优化 + 图片尺寸")
    r = runner.optimize_file(xsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(content=True, image_size=True)), OUT)
    check(not os.path.exists(r.target + ".tmp"), "无残留临时文件")
    from openpyxl import load_workbook

    wb = load_workbook(r.target)
    c3 = wb["清单"]["C3"].value
    c4 = wb["清单"]["C4"].value
    print("   C3 =", repr(c3))
    print("   C4 =", repr(c4))
    check("众所周知" not in (c3 or ""), "Excel 赘语已删除")
    check("设备" in (c3 or "") and "性能" in (c3 or ""), "Excel 关键信息保留")
    ws = wb["清单"]
    check(wb["清单"]["B3"].value == "服务器", "非目标单元格内容不变")
    imgs = getattr(ws, "_images", [])
    check(len(imgs) == 1, "图片保留")
    if imgs:
        ext = getattr(imgs[0].anchor, "ext", None)
        cx = int(getattr(ext, "cx", 0) or 0)
        cy = int(getattr(ext, "cy", 0) or 0)
        print("   anchor.ext(EMU)：", cx, cy,
              "= %.1f×%.1fcm" % (cx / 360000.0, cy / 360000.0))
        check(0 < cx <= 18 * 360000 and 0 < cy <= 12 * 360000,
              "Excel 图片已缩小进 18×12cm 打印区（anchor.ext 已生效）")
    wb.close()

    # ---- 场景 6：排版总开关关闭 ----
    print("\n[6] 「自动排版」总开关关闭，但其余版式开关都开着")
    r = runner.optimize_file(wsrc, tpl, OptimizeOptions(
        "auto", ToggleOptions(layout=False, image_size=False)), OUT)
    check(fp(r.target, "word") == fp(wsrc, "word"), "文字零改动")
    check(os.path.exists(r.target), "仍能正常输出")

    print(f"\n合计：通过 {ok} / 失败 {bad}")
    return 0 if bad == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
